"""Human-labeled per-lesson TalkMoves rates; no model predictions are used.

Run after fetch_talkmoves.sh with numpy, matplotlib, and openpyxl installed.
The upstream notebook gives valid Student Tag labels precedence over Teacher Tag.
Each native workbook is one lesson/lesson portion; zero-teacher books are excluded.
"""

from __future__ import annotations

import hashlib
import json
import re
from collections import Counter

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook

from talkmoves_common import LABELS, MODELS, MOVE_IDS, MOVE_LABELS, RAW, RESULTS, SEED, sha256, write_json

NATIVE_TO_MODEL = {1: 0, 2: 1, 3: 2, 4: 3, 5: 4, 8: 5, 9: 6}
BOOTSTRAPS = 2000


def native_code(value):
    if value is None:
        return None
    match = re.match(r"^\s*(\d+)(?:\s*[-.]|\s*$)", str(value))
    return int(match.group(1)) if match else None


def nonblank(value):
    return value is not None and bool(str(value).strip())


def profile_lessons():
    lessons = []
    counts = Counter()
    unknown_teacher_codes = Counter()
    native_labels = Counter()
    subsets = Counter()
    header_variants = Counter()
    checksum_counts = Counter()
    for subset in ("Subset 1", "Subset 2"):
        for path in sorted((RAW / "data" / subset).glob("*.xlsx")):
            if path.name.startswith("~$"):
                continue
            subsets[subset] += 1
            checksum_counts[sha256(path)] += 1
            workbook = load_workbook(path, read_only=True, data_only=True)
            extra_sheets = workbook.sheetnames[1:]
            if any(name != "Tags" for name in extra_sheets):
                raise ValueError(f"Unexpected additional native sheet: {path.name}")
            counts["auxiliary_Tags_sheets_ignored"] += len(extra_sheets)
            rows = workbook.worksheets[0].iter_rows(values_only=True)
            header = tuple(next(rows))
            columns = {str(value).strip(): index for index, value in enumerate(header)}
            for key in ("Sentence", "Teacher Tag", "Student Tag"):
                if key not in columns:
                    raise ValueError(f"Missing {key} in native workbook: {path.name}")
            header_variants[json.dumps(header, default=str)] += 1
            lesson_labels = Counter()
            for row in rows:
                counts["native_rows"] += 1
                def value(key):
                    index = columns.get(key)
                    return row[index] if index is not None and index < len(row) else None
                teacher = native_code(value("Teacher Tag"))
                student = native_code(value("Student Tag"))
                if teacher is None:
                    counts["missing_or_unparseable_teacher_tag_rows"] += 1
                    continue
                if teacher not in NATIVE_TO_MODEL:
                    unknown_teacher_codes[str(teacher)] += 1
                    continue
                if not nonblank(value("Sentence")):
                    counts["blank_sentence_valid_teacher_tag_rows"] += 1
                    continue
                if student in range(1, 6):
                    counts["dual_valid_student_teacher_rows_excluded"] += 1
                    continue
                label = NATIVE_TO_MODEL[teacher]
                lesson_labels[label] += 1
                native_labels[teacher] += 1
                if value("Speaker") != "T":
                    counts["teacher_rows_lost_if_literal_speaker_T_used"] += 1
            workbook.close()
            # Hash identifiers preserve inspectable lesson grain without exporting names or text.
            source_path = str(path.relative_to(RAW))
            lessons.append({"lesson_id": hashlib.sha256(source_path.encode()).hexdigest()[:16],
                            "subset": subset, "source_sha256": sha256(path),
                            "teacher_sentences": sum(lesson_labels.values()),
                            "labels": {str(label): lesson_labels[label] for label in LABELS}})
    eligible = [lesson for lesson in lessons if lesson["teacher_sentences"]]
    if not eligible:
        raise ValueError("No eligible native teacher lessons found.")
    return lessons, eligible, {
        "native_files": len(lessons), "subsets": dict(subsets), "eligible_lessons": len(eligible),
        "excluded_zero_teacher_lessons": len(lessons) - len(eligible),
        "teacher_sentences": sum(lesson["teacher_sentences"] for lesson in eligible),
        "teacher_native_label_counts": {str(key): value for key, value in sorted(native_labels.items())},
        "unknown_or_non_teacher_native_codes_excluded": dict(sorted(unknown_teacher_codes.items())),
        "identical_native_workbook_excess_files": sum(value - 1 for value in checksum_counts.values()),
        "header_variants": dict(header_variants), **dict(counts),
    }


def plot_histogram(move, rates, summary):
    fig, ax = plt.subplots(figsize=(10, 5.5), constrained_layout=True)
    upper = max(float(rates.max()), 1)
    bins = np.linspace(0, upper, 25)
    ax.axvspan(summary["p25"], summary["p75"], color="#e7ebf0", label="Middle 50% of lessons", zorder=0)
    ax.hist(rates, bins=bins, color="#41678d", edgecolor="white", linewidth=0.6, label="Human-labeled lessons", zorder=2)
    ax.axvline(summary["median"], color="#222b37", linewidth=1.8, label="Median lesson rate", zorder=3)
    ax.set_xlim(0, upper * 1.02)
    ax.set_ylim(bottom=0)
    ax.set_xlabel(f"Human-labeled {move.replace('_', ' ')} moves per 100 teacher sentences")
    ax.set_ylabel("Number of lessons")
    ax.set_title(f"{MOVE_LABELS[move]} — classroom reference distribution\nTalkMoves native lessons (n={len(rates):,}); human annotations, not model predictions", pad=14)
    ci = summary["ci95_median"]
    ax.text(0.98, 0.95, f"Median {summary['median']:.2f} per 100\nIQR {summary['p25']:.2f}–{summary['p75']:.2f}\n95% bootstrap CI for median {ci[0]:.2f}–{ci[1]:.2f}",
            transform=ax.transAxes, ha="right", va="top", fontsize=10,
            bbox={"facecolor": "white", "edgecolor": "#dbe0e6", "pad": 9})
    ax.grid(axis="y", color="#e6e9ed", linewidth=0.6)
    ax.set_axisbelow(True)
    ax.legend(loc="center right", frameon=False, fontsize=9)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    fig.savefig(RESULTS / f"benchmark_{move}.png", dpi=170)
    plt.close(fig)


def main():
    lessons, eligible, quality = profile_lessons()
    moves = list(MOVE_LABELS)
    rates = np.array([[100 * lesson["labels"][str(MOVE_IDS[move])] / lesson["teacher_sentences"] for move in moves] for lesson in eligible])
    # One common lesson-level resampling index matrix preserves the lesson unit.
    resamples = np.random.default_rng(SEED).integers(0, len(eligible), size=(BOOTSTRAPS, len(eligible)))
    boot_medians = np.median(rates[resamples], axis=1)
    benchmarks = {}
    for index, move in enumerate(moves):
        benchmarks[move] = {"median": float(np.median(rates[:, index])),
                            "p25": float(np.percentile(rates[:, index], 25)),
                            "p75": float(np.percentile(rates[:, index], 75)),
                            "ci95_median": np.percentile(boot_medians[:, index], [2.5, 97.5]).tolist(),
                            "lessons": len(eligible)}
        plot_histogram(move, rates[:, index], benchmarks[move])
    manifest = json.loads((RESULTS / "source_manifest.json").read_text())
    metadata = {"source": "Human labels from native data/Subset 1 and data/Subset 2 workbooks, one lesson or lesson portion per file.",
                "provenance": {"revision": manifest["revision"], "repository": manifest["repository"], "license": manifest["license"]},
                "unit": "move count per 100 valid nonblank human-labeled teacher sentences within each lesson",
                "aggregation": "Unweighted distribution across eligible lessons; all seven valid teacher classes in denominator.",
                "teacher_selection": "Student Tag native codes 1–5 take precedence. Otherwise use Teacher Tag codes 1,2,3,4,5,8,9, irrespective of Speaker value. Exclude blank sentences, missing tags, Context7, and undocumented codes6/128.",
                "bootstrap": {"unit": "lesson", "resamples": BOOTSTRAPS, "seed": SEED, "interval": "Percentile 2.5 and 97.5 of bootstrap median", "quantile_method": "numpy linear"},
                "data_quality": quality,
                "limitations": ["Benchmark native workbooks and official model TSVs have distinct extraction histories and need not yield identical totals.",
                                "All eligible native lessons are used; the illustrative 504 count is not forced.",
                                "Descriptive K-12 mathematics reference rates are not causal targets or college teaching-quality scores."]}
    write_json(MODELS / "talkmoves_benchmarks.json", benchmarks)
    write_json(RESULTS / "benchmark_summary.json", {**metadata, "moves": benchmarks})
    write_json(RESULTS / "lesson_rates.json", {"metadata": metadata, "lessons": lessons})
    print(json.dumps({"quality": quality, "benchmarks": benchmarks}), flush=True)


if __name__ == "__main__":
    main()
